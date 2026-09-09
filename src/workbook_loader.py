from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import logging
from pathlib import Path
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


@dataclass
class WorkbookPair:
    values: Workbook
    formulas: Workbook

    def close(self) -> None:
        self.values.close()
        self.formulas.close()


def load_workbook_pair(path: Path) -> WorkbookPair:
    try:
        return _load_pair(path)
    except TypeError as exc:
        if str(exc) != "Nested.from_tree() missing 1 required positional argument: 'node'":
            raise
        # openpyxl 3.1.5 cannot deserialize some pivot cache definitions.
        # GEN150 needs worksheet cells, not editable PivotTable objects.
        # Disconnect only those objects in an in-memory ingestion copy; keep
        # worksheet XML (including cached results), styles and links untouched.
        with _without_pivot_relationships(path) as source:
            logging.getLogger("gen150").warning(
                "Pivot cache compatibility fallback for %s: loading cells without "
                "PivotTable objects; original workbook is unchanged", path.name,
            )
            return _load_pair(source)


def _load_pair(source: Path | BytesIO) -> WorkbookPair:
    values = load_workbook(source, data_only=True, read_only=False, keep_links=True)
    try:
        if isinstance(source, BytesIO):
            source.seek(0)
        formulas = load_workbook(source, data_only=False, read_only=False, keep_links=True)
    except Exception:
        values.close()
        raise
    return WorkbookPair(values, formulas)


def _without_pivot_relationships(path: Path) -> BytesIO:
    """Create an ingestion-only copy; never save this copy over a workbook."""
    result = BytesIO()
    removed = 0
    with ZipFile(path) as source, ZipFile(result, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                # Remove only the cache declaration, preserving all other XML
                # bytes and namespace prefixes (notably extension metadata).
                data = re.sub(
                    rb'<(?P<tag>(?:[A-Za-z_][\w.-]*:)?pivotCaches)\b[^>]*'
                    rb'(?:/>|>.*?</(?P=tag)\s*>)',
                    b"", data, flags=re.DOTALL,
                )
            if entry.filename.startswith("xl/worksheets/_rels/") and entry.filename.endswith(".rels"):
                relationships = ElementTree.fromstring(data)
                pivots = [node for node in relationships if node.get("Type", "").endswith("/pivotTable")]
                if pivots:
                    for node in pivots:
                        relationships.remove(node)
                    removed += len(pivots)
                    data = ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, data)
    if not removed:
        result.close()
        raise ValueError("Pivot cache fallback found no worksheet PivotTable relationships")
    result.seek(0)
    return result
