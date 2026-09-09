from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string

from ..date_utils import extract_reporting_date, normalize_reporting_date, parse_period
from ..numeric_utils import percentage_decimal, scaled_decimal, to_decimal
from ..text_utils import canonical_text, clean_text, contains_total
from .base_parser import BaseParser, ParseResult


@dataclass(frozen=True)
class CategoryBlock:
    label: str
    first_column: int
    last_column: int


class ApsfCompetitorSummaryParser(BaseParser):
    HEADER = (
        "date_reporting",
        "section",
        "organisme",
        "categorie",
        "periode_reference",
        "periode_comparaison",
        "valeur_reference_mdh",
        "valeur_comparaison_mdh",
        "evolution_pct",
        "marche_reference_mdh",
        "marche_comparaison_mdh",
        "evolution_marche_pct",
        "ratio_reference_pct",
        "ratio_comparaison_pct",
        "evolution_ratio_points",
        "marche_ratio_reference_pct",
        "marche_ratio_comparaison_pct",
        "evolution_marche_ratio_points",
    )
    STAGING_HEADER = (
        "date_reporting",
        "sheet",
        "section",
        "category_label_raw",
        "period_label_raw",
        "row_label_raw",
        "row_type",
        "value",
        "number_format",
        "source_row",
        "source_column",
    )

    RATIO_SECTION = "CES_SUR_ENCOURS_BRUT"

    def process(self) -> ParseResult:
        layout = self.config.get("layout", "summary")
        if layout == "staging":
            return self._process_staging()
        if layout != "summary":
            raise ValueError(f"Unsupported competitor layout: {layout!r}")
        return self._process_summary()

    def _process_summary(self) -> ParseResult:
        title_area = self.config.get(
            "title_area",
            {"first_row": 1, "last_row": 8, "first_column": 1, "last_column": 12},
        )
        date_reporting = extract_reporting_date(self.values_in_area(title_area))
        label_column = column_index_from_string(self.config["label_column"])
        market_label = canonical_text(self.config.get("market_label", "Marché"))
        market_share_label = canonical_text(
            self.config.get("market_share_label", "Part de marché")
        )
        rows: list[dict[str, Any]] = []

        for section, section_config in self.config["sections"].items():
            start = int(section_config["data_start_row"])
            end = int(section_config["data_end_row"])
            for category, columns in self.config["category_blocks"].items():
                reference_column, comparison_column, evolution_column = (
                    column_index_from_string(column) for column in columns
                )
                period_row = int(section_config["period_row"])
                reference_period = parse_period(self.value(period_row, reference_column))
                comparison_period = parse_period(self.value(period_row, comparison_column))
                market_values: tuple[Any, Any, Any] | None = None
                organization_values: list[tuple[str, tuple[Any, Any, Any]]] = []

                for row_number in range(start, end + 1):
                    label = clean_text(self.value(row_number, label_column))
                    canonical = canonical_text(label)
                    if not canonical or contains_total(label) or market_share_label in canonical:
                        continue
                    values = self._read_triplet(
                        row_number,
                        reference_column,
                        comparison_column,
                        evolution_column,
                        section == self.RATIO_SECTION,
                    )
                    if canonical == market_label:
                        market_values = values
                        continue
                    if any(value not in (None, "") for value in values):
                        organization_values.append((label, values))

                if organization_values and market_values is None and self.config.get(
                    "require_market_context", True
                ):
                    raise ValueError(
                        f"Market row missing for section={section}, category={category}"
                    )
                market_values = market_values or (None, None, None)
                for organisme, values in organization_values:
                    rows.append(
                        self._build_row(
                            date_reporting,
                            section,
                            organisme,
                            category,
                            reference_period,
                            comparison_period,
                            values,
                            market_values,
                        )
                    )

        return ParseResult(self.HEADER, rows, self.warnings)

    def _read_triplet(
        self,
        row: int,
        reference_column: int,
        comparison_column: int,
        evolution_column: int,
        is_ratio: bool,
    ) -> tuple[Any, Any, Any]:
        cells = [
            self.ws.cell(row=row, column=reference_column),
            self.ws.cell(row=row, column=comparison_column),
            self.ws.cell(row=row, column=evolution_column),
        ]
        values = [self.value(row, cell.column) for cell in cells]
        if is_ratio:
            return tuple(
                percentage_decimal(value, cell.number_format)
                for value, cell in zip(values, cells)
            )
        return (
            scaled_decimal(values[0], cells[0].number_format),
            scaled_decimal(values[1], cells[1].number_format),
            percentage_decimal(values[2], cells[2].number_format),
        )

    def _build_row(
        self,
        date_reporting: str,
        section: str,
        organisme: str,
        category: str,
        reference_period: str,
        comparison_period: str,
        values: tuple[Any, Any, Any],
        market_values: tuple[Any, Any, Any],
    ) -> dict[str, Any]:
        row = {field: None for field in self.HEADER}
        row.update(
            {
                "date_reporting": date_reporting,
                "section": section,
                "organisme": organisme,
                "categorie": category,
                "periode_reference": reference_period,
                "periode_comparaison": comparison_period,
            }
        )
        if section == self.RATIO_SECTION:
            row.update(
                {
                    "ratio_reference_pct": values[0],
                    "ratio_comparaison_pct": values[1],
                    "evolution_ratio_points": values[2],
                    "marche_ratio_reference_pct": market_values[0],
                    "marche_ratio_comparaison_pct": market_values[1],
                    "evolution_marche_ratio_points": market_values[2],
                }
            )
        else:
            row.update(
                {
                    "valeur_reference_mdh": values[0],
                    "valeur_comparaison_mdh": values[1],
                    "evolution_pct": values[2],
                    "marche_reference_mdh": market_values[0],
                    "marche_comparaison_mdh": market_values[1],
                    "evolution_marche_pct": market_values[2],
                }
            )
        return row

    def _process_staging(self) -> ParseResult:
        date_reporting = self._reporting_date()
        label_column = column_index_from_string(self.config["label_column"])
        market_label = canonical_text(self.config.get("market_label", "Marché"))
        market_share_label = canonical_text(
            self.config.get("market_share_label", "Part de marché")
        )
        rows: list[dict[str, Any]] = []

        for section, section_config in self.config["sections"].items():
            period_row = int(section_config["period_row"])
            blocks = self._category_blocks(int(section_config["category_row"]))
            for block in blocks:
                for column in range(block.first_column, block.last_column + 1):
                    period_label = self._readable_text(self.value(period_row, column))
                    for row_number in range(
                        int(section_config["data_start_row"]),
                        int(section_config["data_end_row"]) + 1,
                    ):
                        row_label = clean_text(self.value(row_number, label_column))
                        canonical_label = canonical_text(row_label)
                        if canonical_label == market_label:
                            row_type = "marche"
                        elif canonical_label == market_share_label:
                            row_type = "part_marche"
                        else:
                            row_type = "organisme"
                        cell = self.ws.cell(row=row_number, column=column)
                        rows.append(
                            {
                                "date_reporting": date_reporting,
                                "sheet": self.ws.title,
                                "section": section,
                                "category_label_raw": block.label,
                                "period_label_raw": period_label,
                                "row_label_raw": row_label,
                                "row_type": row_type,
                                "value": to_decimal(self.value(row_number, column)),
                                "number_format": cell.number_format,
                                "source_row": row_number,
                                "source_column": get_column_letter(column),
                            }
                        )

        return ParseResult(self.STAGING_HEADER, rows, self.warnings)

    def _reporting_date(self) -> str:
        configured = self.config.get("date_reporting")
        if configured not in (None, ""):
            return normalize_reporting_date(configured)
        title_area = self.config.get(
            "title_area",
            {"first_row": 1, "last_row": 8, "first_column": 1, "last_column": 12},
        )
        return extract_reporting_date(self.values_in_area(title_area))

    def _category_blocks(self, category_row: int) -> list[CategoryBlock]:
        limits = self.config.get("limits", {})
        first_allowed = int(limits.get("first_column", 1))
        last_allowed = int(limits.get("last_column", self.ws.max_column))
        blocks: list[CategoryBlock] = []
        for merged_range in self.ws.merged_cells.ranges:
            if merged_range.min_row != category_row or merged_range.max_row != category_row:
                continue
            if merged_range.min_col < first_allowed or merged_range.max_col > last_allowed:
                continue
            label = clean_text(self.value(category_row, merged_range.min_col))
            if not label:
                continue
            blocks.append(
                CategoryBlock(label, merged_range.min_col, merged_range.max_col)
            )
        blocks.sort(key=lambda block: block.first_column)
        if not blocks:
            raise ValueError(
                f"No merged category block found on {self.ws.title!r} row {category_row}"
            )
        return blocks

    @staticmethod
    def _readable_text(value: object) -> str:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return clean_text(value)
