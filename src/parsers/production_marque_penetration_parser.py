from __future__ import annotations

from datetime import date, datetime
from typing import Any

from openpyxl.utils.cell import column_index_from_string

from ..numeric_utils import percentage_decimal, to_decimal
from ..text_utils import canonical_text
from .base_parser import BaseParser, ParseResult


FRENCH_MONTH_NAMES = (
    "", "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
)
FRENCH_MONTH_ABBREVIATIONS = (
    "", "janv", "févr", "mars", "avr", "mai", "juin",
    "juil", "août", "sept", "oct", "nov", "déc",
)


class ProductionMarquePenetrationParser(BaseParser):
    """Read MARQUE or VILLE penetration tables.

    The source sheet 'ville' also owns C5, the reference date used by formulas
    in the other sheets. Keep that source cell even though it is not exported.
    """

    HEADER = (
        "marque", "ville", "periodicite", "dimension",
        "axe", "sous_axe", "valeur",
    )

    def process(self) -> ParseResult:
        label_column = column_index_from_string(self.config["label_column"])
        dimension = self.config["dimension"]
        if dimension not in {"MARQUE", "VILLE"}:
            raise ValueError(f"Invalid dimension: {dimension!r}")
        ville = self.config["ville"]
        rows: list[dict[str, Any]] = []
        tables = self.config.get("tables") or [{}]
        for table in tables:
            self._append_table(rows, table, label_column, dimension, ville)
        return ParseResult(self.HEADER, rows, self.warnings)

    def _append_table(
        self,
        rows: list[dict[str, Any]],
        table: dict[str, Any],
        label_column: int,
        dimension: str,
        ville: Any,
    ) -> None:
        periodicite = table.get("periodicite", self.config.get("periodicite"))
        if periodicite not in {"MENSUEL", "ANNUEL"}:
            raise ValueError(f"Invalid periodicite: {periodicite!r}")
        axis_row = int(table.get("axis_row", self.config["axis_row"]))
        header_row = int(table.get("header_row", self.config["header_row"]))
        data_start_row = int(table.get("data_start_row", self.config["data_start_row"]))
        if canonical_text(self.value(header_row, label_column)) != dimension:
            raise ValueError(f"Expected {dimension} at the configured header row")
        measures: list[tuple[Any, Any, int, bool]] = []
        axis_groups = table.get("axis_groups", self.config.get("axis_groups", []))
        for group in axis_groups:
            axis_column = column_index_from_string(group["axis_column"])
            axis = self._axis_label(self.value(axis_row, axis_column))
            if axis in (None, ""):
                raise ValueError(f"Missing penetration axis at column {group['axis_column']}")
            is_percentage = canonical_text(axis) in {"TAUX_DE_PENETRATION", "VARIATION_TPR"}
            for configured_column in group["value_columns"]:
                column = column_index_from_string(configured_column)
                sub_axis = self._axis_label(
                    self.value(header_row, column),
                    self.ws.cell(header_row, column).number_format,
                )
                if sub_axis in (None, ""):
                    raise ValueError(f"Missing penetration sub-axis at column {configured_column}")
                measures.append((axis, sub_axis, column, is_percentage))
        if not measures:
            raise ValueError("Penetration axis_groups must define at least one value column")
        # Men/An share F7 (=C7-1) and I7 (=C5-365), but their VLOOKUPs
        # reference distinct pivot blocks. In the supplied global source,
        # E:G selects March 2026, Q:S January-March 2026; I:K selects April
        # 2025, U:W January-April 2025. The global annual view is cumulative.
        # This does not confirm every city filter: periodicite is explicit
        # metadata, not a date interval inferred from these ambiguous labels.
        end = int(table.get("data_end_row", self.config.get("data_end_row", self.ws.max_row)))
        for row_number in range(data_start_row, end + 1):
            marque = self.value(row_number, label_column)
            if marque in (None, "") or canonical_text(marque) == "TOTAL_GENERAL":
                break
            for axis, sub_axis, column, is_percentage in measures:
                value = self.value(row_number, column)
                converted = (
                    percentage_decimal(value, self.ws.cell(row_number, column).number_format)
                    if is_percentage else to_decimal(value)
                )
                rows.append({
                    "marque": marque if dimension == "MARQUE" else None,
                    "ville": marque if dimension == "VILLE" else ville,
                    "periodicite": periodicite,
                    "dimension": dimension,
                    "axe": axis,
                    "sous_axe": sub_axis,
                    "valeur": converted,
                })

    @staticmethod
    def _axis_label(value: Any, number_format: str | None = None) -> Any:
        if isinstance(value, (datetime, date)):
            if number_format and "TPR" in number_format.upper():
                month = FRENCH_MONTH_ABBREVIATIONS[value.month]
                return f"TPR {month} {value.year}"
            return f"{FRENCH_MONTH_NAMES[value.month]} {value.year}"
        return value
