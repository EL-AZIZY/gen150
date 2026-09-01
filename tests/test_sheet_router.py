import logging
import unittest

from openpyxl import Workbook

from src.sheet_router import SheetRouter, find_sheet_name, normalize_sheet_name


class SheetRouterTests(unittest.TestCase):
    def test_normalization_and_find_preserve_real_name(self):
        workbook = Workbook()
        workbook.active.title = "RCI "
        self.assertEqual(normalize_sheet_name("  rci  "), "rci")
        self.assertEqual(find_sheet_name(workbook, " rCi "), "RCI ")

    def test_only_declared_visible_sheet_is_routed_once(self):
        workbook = Workbook()
        workbook.active.title = "RCI "
        workbook.create_sheet("Hidden").sheet_state = "hidden"
        workbook.create_sheet("VeryHidden").sheet_state = "veryHidden"
        workbook.create_sheet("Unconfigured")
        families = {
            "family": {
                "enabled": True,
                "family_sheets": ["RCI", "Hidden", "VeryHidden"],
            }
        }
        structures = {
            "sheets": {
                "RCI": {"family": "family"},
                "Hidden": {"family": "family"},
                "VeryHidden": {"family": "family"},
            }
        }
        routed = list(SheetRouter(workbook, families, structures, logging.getLogger()).route())
        self.assertEqual([item.real_name for item in routed], ["RCI "])

    def test_duplicate_normalized_sheet_names_are_rejected(self):
        workbook = Workbook()
        workbook.active.title = "RCI"
        workbook.create_sheet(" rci ")
        with self.assertRaises(ValueError):
            SheetRouter(workbook, {}, {}, logging.getLogger())


if __name__ == "__main__":
    unittest.main()
